from __future__ import annotations

from .i18n import tr, fmt, join_text
from datetime import datetime, timezone
from decimal import Decimal
import os
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .domain import CellResult, FieldRegion, JobResult, PageResult, TableTemplate
from .ocr import is_simple_number


def _excel_value(text: str, *, numeric: bool = True) -> Any:
    value = text.strip()
    if not value:
        return None
    if numeric and is_simple_number(value):
        normalized = value.replace("−", "-").replace(",", ".")
        number = float(normalized)
        return int(number) if number.is_integer() and "." not in normalized else number
    return value


def _field_result(page: PageResult, region: FieldRegion) -> str:
    result = next((item for item in page.fields if item.region_id == region.id), None)
    return result.final_text if result else ""


def _group_value(page: PageResult, regions: list[FieldRegion], column: int) -> str:
    for region in regions:
        if region.export_mode != "column_group":
            continue
        if region.column_start is not None and region.column_end is not None and region.column_start <= column <= region.column_end:
            return _field_result(page, region)
    return ""


def _cell(page: PageResult, row: int, column: int) -> CellResult | None:
    return page.cell(row, column)


def _separator_evidence(raw_text: str, final_text: str, flags: list[str]) -> str:
    if "," in raw_text:
        return "comma → canonical dot"
    if "." in raw_text:
        return "dot"
    if "separator_reclassified_from_one" in flags:
        return "OCR 1 → proposed dot from template"
    if "separator_inferred_from_rule" in flags:
        return "missing → proposed dot from template"
    if "." in final_text:
        return "dot not present in primary OCR"
    return ""


def _layout_sheet_title(page_index: int) -> str:
    return tr('Исходная таблица') if page_index == 0 else tr('Исходная таблица {p0}', p0=page_index + 1)


def _is_header_cell(template: TableTemplate, row: int, column: int) -> bool:
    return template.is_header_cell(row, column)


def _cell_value(result: CellResult, template: TableTemplate) -> tuple[Any, str]:
    rule, _ = template.cell_constraints(result.row, result.column)
    numeric = bool(rule and rule.value_format in {"numeric", "integer", "complex_numeric"}
                   and not template.is_label_cell(result.column))
    value = _excel_value(result.final_text, numeric=numeric)
    number_format = "General"
    if numeric and isinstance(value, (float, int)):
        places = None if result.status in {"confirmed", "corrected"} else rule.decimal_places
        if places is None:
            normalized = result.final_text.strip().replace(",", ".").replace("−", "-")
            places = max(0, -Decimal(normalized).as_tuple().exponent)
        number_format = "0." + "0" * places if places else "0"
    return value, number_format


def _add_layout_sheet(workbook: Workbook, page: PageResult, job: JobResult) -> None:
    """Add the reviewed cell matrix in the same row/column layout as the scan."""
    sheet = workbook.create_sheet(_layout_sheet_title(page.page_index))
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"A{job.template.header_rows + 1}" if job.template.header_rows else None

    thin = Side(style="thin", color="AEB7C4")
    grid_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E8EEF8")
    label_fill = PatternFill("solid", fgColor="F4F6F8")
    excluded_fill = PatternFill("solid", fgColor="ECEFF3")

    for row in range(job.template.rows):
        excluded = row in page.excluded_rows
        for column in range(job.template.columns):
            result = page.cell(row, column)
            is_excluded_value = bool(
                result and (result.status == "excluded" or page.is_excluded(row, column, job.template))
            )
            value, number_format = (None, "General") if result is None or is_excluded_value else _cell_value(result, job.template)
            cell = sheet.cell(row=row + 1, column=column + 1, value=value)
            cell.number_format = number_format
            cell.border = grid_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if _is_header_cell(job.template, row, column):
                cell.fill = header_fill
                cell.font = Font(bold=True, color="1F2937")
            elif column < job.template.row_label_columns:
                cell.fill = label_fill
                cell.font = Font(bold=True, color="1F2937")
            elif is_excluded_value:
                cell.fill = excluded_fill

    column_spans = [
        job.template.column_guides[index + 1] - job.template.column_guides[index]
        for index in range(job.template.columns)
    ]
    median_column = sorted(column_spans)[len(column_spans) // 2] if column_spans else 1
    for index, span in enumerate(column_spans, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = max(8, min(28, 12 * span / max(median_column, 1e-9)))

    row_spans = [
        job.template.row_guides[index + 1] - job.template.row_guides[index]
        for index in range(job.template.rows)
    ]
    median_row = sorted(row_spans)[len(row_spans) // 2] if row_spans else 1
    for index, span in enumerate(row_spans, start=1):
        sheet.row_dimensions[index].height = max(18, min(42, 21 * span / max(median_row, 1e-9)))

    sheet.print_area = f"A1:{get_column_letter(job.template.columns)}{job.template.rows}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def export_job(job: JobResult, target: str | Path, *, mode: str = "extended") -> Path:
    """Export reviewed results; compact contains source layouts and fields."""
    if mode not in {"compact", "extended"}:
        raise ValueError(tr('Неизвестный режим экспорта'))
    if not job.pages:
        raise ValueError(tr('Нет страниц для экспорта'))
    job.template.validate_value_rules()
    if job.unresolved_count:
        raise ValueError(tr('{p0} values still need review', p0=job.unresolved_count))
    for page in job.pages:
        for region in job.template.fields:
            reviewed = any(f.region_id == region.id and f.status in {"confirmed", "corrected"} for f in page.fields)
            if not reviewed and region.hard_errors(_field_result(page, region)):
                raise ValueError(tr('Поле «{p0}» не соответствует правилу: {p1}', p0=region.name, p1=region.constraints().summary()))
        for cell in page.cells:
            if cell.status not in {"excluded", "confirmed", "corrected"} and not page.is_excluded(cell.row, cell.column, job.template):
                rule, name = job.template.cell_constraints(cell.row, cell.column)
                if rule and rule.hard_errors(cell.final_text):
                    raise ValueError(tr('R{p0}C{p1}: значение не соответствует правилу «{p2}»', p0=cell.row + 1, p1=cell.column + 1, p2=name))

    target_path = Path(target)
    target_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    if mode == "compact":
        workbook.remove(workbook.active)
        for page in job.pages:
            _add_layout_sheet(workbook, page, job)
        if any(page.fields for page in job.pages):
            fields_sheet = workbook.create_sheet('Fields')
            fields_sheet.append(['Page', 'Field', 'Value'])
            for page in job.pages:
                for field in page.fields:
                    # Metadata is text: identifiers, dates and leading zeros
                    # must survive exactly as reviewed, including formula-like text.
                    fields_sheet.append([page.page_index + 1, field.name, field.final_text])
            fields_sheet.freeze_panes = 'A2'
            fields_sheet.auto_filter.ref = fields_sheet.dimensions
            for cell in fields_sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='1D4ED8')
            for column, width in [('A', 10), ('B', 28), ('C', 60)]:
                fields_sheet.column_dimensions[column].width = width
            for row in fields_sheet.iter_rows(min_row=2):
                row[2].alignment = Alignment(wrap_text=True, vertical='top')
        return _save_workbook(workbook, target_path)
    data_sheet = workbook.active
    data_sheet.title = "Data"

    repeated_regions = [region for region in job.template.fields if region.export_mode == "repeat"]
    group_regions = [region for region in job.template.fields if region.export_mode == "column_group"]
    repeated_names = list(dict.fromkeys(region.name for region in repeated_regions))
    headers = ["Source file", "Page", *repeated_names]
    if group_regions:
        headers.append("Side")
    headers.extend(["Source row", "Measurement", "Value", "Review status", "Confidence"])
    data_sheet.append(headers)

    for page in job.pages:
        repeated_values: dict[str, str] = {}
        for name in repeated_names:
            matching = next((region for region in repeated_regions if region.name == name), None)
            repeated_values[name] = _field_result(page, matching) if matching else ""

        for row in range(job.template.rows):
            data_columns = [column for column in range(job.template.row_label_columns, job.template.columns)
                            if job.template.column_rules[column].role == "data"
                            and not _is_header_cell(job.template, row, column)]
            if not data_columns:
                continue
            row_label_parts = []
            for column in range(job.template.row_label_columns):
                result = _cell(page, row, column)
                if result and result.final_text:
                    row_label_parts.append(result.final_text)
            earlier_headers = sum(all(_is_header_cell(job.template, previous, column) for column in data_columns)
                                  for previous in range(row))
            row_label = " ".join(row_label_parts) or str(row + 1 - earlier_headers)

            if row in page.excluded_rows:
                record = [job.source_name, page.page_index + 1, *[repeated_values[name] for name in repeated_names]]
                if group_regions:
                    record.append("")
                record.extend([row_label, "", None, "excluded", 1.0])
                data_sheet.append(record)
                continue

            for column in data_columns:
                rule = job.template.column_rules[column]
                if rule.role != "data":
                    continue
                result = _cell(page, row, column)
                if result is None or not result.final_text or result.status == 'excluded' or page.is_excluded(row, column, job.template):
                    continue
                header_parts = []
                for header_row in range(job.template.header_rows):
                    if not _is_header_cell(job.template, header_row, column):
                        continue
                    header_cell = _cell(page, header_row, column)
                    if header_cell and header_cell.final_text:
                        header_parts.append(header_cell.final_text)
                measurement = rule.name if not rule.name.startswith("Column ") else " / ".join(header_parts) or str(column + 1)
                record = [job.source_name, page.page_index + 1, *[repeated_values[name] for name in repeated_names]]
                if group_regions:
                    record.append(_group_value(page, group_regions, column))
                value, number_format = _cell_value(result, job.template)
                record.extend([row_label, measurement, value, result.status, round(result.confidence, 4)])
                data_sheet.append(record)
                data_sheet.cell(data_sheet.max_row, headers.index("Value") + 1).number_format = number_format

    for page in job.pages:
        _add_layout_sheet(workbook, page, job)

    audit_sheet = workbook.create_sheet("Audit")
    audit_headers = [
        "Source file", "Page", "Kind", "Address", "Raw value", "Final value", "Confidence",
        "Flags", "Status", "Template", "Template version", "Model version", "Exported at UTC",
        "Decimal separator evidence", "Applied value rule", "Alternative readings",
        "Writer style suggestion", "Writer adaptation evidence",
    ]
    audit_sheet.append(audit_headers)
    exported_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for page in job.pages:
        for item in page.fields:
            audit_sheet.append([
                job.source_name, page.page_index + 1, "field", item.name, item.raw_text, item.final_text,
                round(item.confidence, 4), ", ".join(item.flags), item.status, job.template.name,
                job.template.template_version, job.model_version, exported_at,
                _separator_evidence(item.raw_text, item.final_text, item.flags), "", item.alternatives, "", "",
            ])
        for item in page.cells:
            audit_sheet.append([
                job.source_name, page.page_index + 1, "cell", f"R{item.row + 1}C{item.column + 1}",
                item.raw_text, item.final_text, round(item.confidence, 4), ", ".join(item.flags), item.status,
                job.template.name, job.template.template_version, job.model_version, exported_at,
                _separator_evidence(item.raw_text, item.final_text, item.flags), item.applied_rule, item.alternatives,
                item.writer_suggestion,
                item.writer_evidence,
            ])

    for sheet in (data_sheet, audit_sheet):
        # OCR and user text are data, never executable workbook formulas.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
            cell.alignment = Alignment(vertical="center")
        for column_cells in sheet.columns:
            width = min(36, max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2))
            sheet.column_dimensions[column_cells[0].column_letter].width = width
        if sheet.max_row > 1:
            table = Table(displayName=f"{sheet.title.replace(' ', '')}Table", ref=sheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
            sheet.add_table(table)
    return _save_workbook(workbook, target_path)


def _save_workbook(workbook: Workbook, target_path: Path) -> Path:
    # All sheets, including the original layout, contain data rather than formulas.
    for sheet in workbook:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    # Keep the previous export intact until the complete replacement is valid.
    # Closing the temporary handle first also permits openpyxl to open it on Windows.
    with NamedTemporaryFile(prefix=f".{target_path.stem}-", suffix=".xlsx", dir=target_path.parent, delete=False) as temporary:
        pending = Path(temporary.name)
    try:
        workbook.save(pending)
        load_workbook(pending, read_only=True, data_only=False).close()
        with ZipFile(pending) as archive:
            broken_member = archive.testzip()
            if broken_member:
                raise ValueError(tr('Повреждён внутренний файл Excel: {p0}', p0=broken_member))
        # Windows requires a writable descriptor for fsync/_commit.
        with pending.open("r+b") as handle:
            os.fsync(handle.fileno())
        os.replace(pending, target_path)
    finally:
        pending.unlink(missing_ok=True)
    return target_path

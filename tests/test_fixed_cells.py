import json
from pathlib import Path
from unittest.mock import Mock

import numpy as np
import pytest
from openpyxl import load_workbook

from tablescan_local.domain import TableTemplate, FixedCell, NormalizedRect, JobResult, FieldRegion
from tablescan_local.ocr import OcrValue
from tablescan_local.pipeline import process_page
from tablescan_local.exporter import export_job
from tablescan_local.storage import LocalStore


def template():
    t = TableTemplate('fixed', 'Fixed test', NormalizedRect(0, 0, 1, 1),
                      [0, .2, .6, 1], [0, .5, 1], header_rows=1, detect_crossed_rows=False)
    t.ensure_column_rules()
    t.column_rules[0].name = 'ID'
    t.column_rules[1].name = 'L1'
    t.set_fixed_value(0, 0, 'ID'); t.set_fixed_value(0, 1, 'L1')
    return t


def test_fixed_cells_skip_crops_ocr_progress_and_keep_real_ids(monkeypatch, tmp_path):
    from tablescan_local import pipeline
    t = template()
    t.fields = [FieldRegion('assay', 'Assay', NormalizedRect(0, 0, .1, .1), source='fixed', fixed_value='Von Frey')]
    engine = Mock()
    engine.recognize_cell.side_effect = [OcrValue(v, .99) for v in ('16', '38.2', '35', '53.6')]
    original = pipeline.cell_crop_bundle
    cropped = []
    def crop(image, template, row, col):
        cropped.append((row, col))
        return original(image, template, row, col)
    monkeypatch.setattr(pipeline, 'cell_crop_bundle', crop)
    progress = []
    page = process_page(np.full((300, 200, 3), 255, np.uint8), 'sample.pdf', 0, t, engine,
                        lambda *args: progress.append(args), tmp_path / 'crops')
    assert cropped == [(1, 0), (1, 1), (2, 0), (2, 1)]
    engine.recognize_region.assert_not_called()
    assert page.cell(0, 1).final_text == 'L1'
    assert page.cell(0, 1).raw_text == '' and not page.cell(0, 1).crop_path
    assert not page.cell(0, 1).needs_review
    assert [page.cell(r, 0).final_text for r in (1, 2)] == ['16', '35']
    assert len(progress) == 4 and progress[-1][:2] == (4, 4)
    assert not page.fields[0].crop_path
    job = JobResult('sample.pdf', t, [page])
    restored = JobResult.from_dict(job.to_dict())
    assert restored.template.fixed_value(0, 1) == 'L1'
    target = export_job(restored, tmp_path / 'result.xlsx', mode='compact')
    workbook = load_workbook(target)
    assert workbook.active['A1'].value == 'ID'
    assert workbook.active['B1'].value == 'L1'
    assert str(workbook.active['A2'].value) == '16'
    assert str(workbook.active['A3'].value) == '35'
    workbook.close()


def test_fixed_empty_overrides_numeric_constraints_and_slow_checks(tmp_path):
    from tablescan_local.slow_mode import eligible_cells, mask_fixed_cells
    t = template(); t.set_fixed_value(1, 1, '')
    engine = Mock(); engine.recognize_cell.return_value = OcrValue('16', .99)
    page = process_page(np.zeros((300, 200, 3), np.uint8), 'sample', 0, t, engine, crop_directory=tmp_path)
    assert t.cell_constraints(1, 1)[0] is None
    assert page.cell(1, 1).final_text == ''
    page.cell(1, 1).flags.append('low_confidence')
    assert page.cell(1, 1) not in eligible_cells(page, t)
    masked = mask_fixed_cells(np.zeros((300, 200, 3), np.uint8), t)
    assert np.all(masked[80:150, 120:180] == 255)
    assert np.all(masked[80:150, 20:80] == 0)


def test_validation_keeps_removed_addresses_visible_and_rejects_duplicates():
    t = template(); t.set_fixed_value(2, 1, 'constant')
    t.resize_grid(2, 2)
    with pytest.raises(ValueError): t.validate_value_rules()
    t.set_fixed_value(2, 1, None)
    t.validate_value_rules()
    t.fixed_cells.append(FixedCell(0, 0, 'duplicate'))
    with pytest.raises(ValueError): t.validate_value_rules()


def test_legacy_templates_still_use_ocr():
    data = template().to_dict(); data.pop('fixed_cells')
    assert TableTemplate.from_dict(data).fixed_value(0, 1) is None


def test_all_fixed_template_does_not_load_ocr_models(monkeypatch):
    from tablescan_local import pipeline, template_fit
    t = template()
    for r in range(t.rows):
        for c in range(t.columns):
            t.set_fixed_value(r, c, str(r * 10 + c))
    monkeypatch.setattr(template_fit, 'fit_document_template', lambda template, images: template)
    engine = Mock(side_effect=AssertionError('OCR models must not load'))
    monkeypatch.setattr(pipeline, 'LocalOcrEngine', engine)
    result = pipeline.process_document([np.full((300, 200, 3), 255, np.uint8)], 'fixed', t)
    engine.assert_not_called()
    assert result.pages[0].cell(2, 1).final_text == '21'


def test_bundled_headers_fixed_but_ids_and_measurements_are_not():
    root = Path(__file__).parents[1] / 'src/tablescan_local/default_templates'
    for path in root.glob('*.json'):
        t = TableTemplate.from_dict(json.loads(path.read_text()))
        t.validate_value_rules()
        assert [t.fixed_value(0, c) for c in range(t.columns)] == [r.name for r in t.column_rules]
        assert all(t.fixed_value(r, c) is None for r in range(1, t.rows) for c in range(t.columns))


def test_existing_untouched_defaults_upgrade_but_edits_remain(tmp_path):
    store = LocalStore(tmp_path); store.install_default_templates()
    templates = store.load_templates()
    for t in templates:
        t.fixed_cells = []; t.schema_version = 3; t.template_version = 2
        store.save_template(t)
    edited = templates[0]; edited.name = 'My protocol'; store.save_template(edited)
    store.install_default_templates()
    after = store.load_templates()
    assert len(after) == 4
    assert next(t for t in after if t.id == edited.id).fixed_cells == []
    assert all(t.fixed_cells for t in after if t.id != edited.id)
    store.install_default_templates()
    assert len(store.load_templates()) == 4


def test_editor_sets_and_removes_constants_and_auto_fills_headers(qtbot):
    from tablescan_local.ui import TablePage
    page = TablePage(mode='template'); qtbot.addWidget(page)
    t = template(); t.fixed_cells = []
    page.set_document(np.full((300, 200, 3), 255, np.uint8), t)
    page.tabs.setCurrentIndex(3)
    page.fixed_panel.fill_headers()
    assert t.fixed_value(0, 1) == 'L1' and t.fixed_value(1, 0) is None
    page.canvas.select_cells({(1, 1)})
    page.fixed_panel.value.setText('same')
    page.fixed_panel.set_values()
    assert t.fixed_value(1, 1) == 'same'
    page.fixed_panel.restore_ocr()
    assert t.fixed_value(1, 1) is None

from zipfile import ZipFile

from openpyxl import load_workbook

from tablescan_local.domain import (
    CellResult,
    ColumnRule,
    FieldRegion,
    FieldResult,
    JobResult,
    NormalizedRect,
    PageResult,
    TableTemplate,
)
from tablescan_local.exporter import export_job


def make_job() -> JobResult:
    template = TableTemplate(
        id="t1",
        name="Plantar",
        table_rect=NormalizedRect(0.1, 0.2, 0.8, 0.6),
        row_guides=[0.2, 0.3, 0.4, 0.5],
        column_guides=[0.1, 0.2, 0.4, 0.6],
        header_rows=1,
        row_label_columns=1,
        fields=[
            FieldRegion("test", "Test", NormalizedRect(0.1, 0.1, 0.2, 0.05)),
            FieldRegion(
                "side", "Side", NormalizedRect(0.3, 0.1, 0.2, 0.05), source="fixed",
                fixed_value="Left Hind", export_mode="column_group", column_start=1, column_end=2,
            ),
        ],
        column_rules=[
            ColumnRule("Animal", role="row_label"),
            ColumnRule("Measure 1"),
            ColumnRule("Measure 2"),
        ],
    )
    texts = [
        ["", "1", "2"],
        ["A1", "12.5", "8.2 ± 0.4"],
        ["A2", "", "15%"],
    ]
    cells = [
        CellResult(row, column, value, value, 0.99, status="confirmed")
        for row, values in enumerate(texts)
        for column, value in enumerate(values)
    ]
    fields = [
        FieldResult("test", "Test", "PLANTAR", "PLANTAR", 0.99, status="confirmed"),
        FieldResult("side", "Side", "Left Hind", "Left Hind", 1.0, status="confirmed"),
    ]
    return JobResult("/tmp/source.pdf", template, [PageResult(0, "/tmp/source.pdf", cells, fields)])


def test_export_creates_normalized_data_and_audit(tmp_path) -> None:
    output = export_job(make_job(), tmp_path / "result.xlsx")
    workbook = load_workbook(output, data_only=True)

    assert workbook.sheetnames == ["Data", "Original table", "Audit"]
    data = workbook["Data"]
    headers = [cell.value for cell in data[1]]
    assert headers == [
        "Source file", "Page", "Test", "Side", "Source row", "Measurement",
        "Value", "Review status", "Confidence",
    ]
    assert data["G2"].value == 12.5
    assert data["G3"].value == "8.2 ± 0.4"
    assert data["D2"].value == "Left Hind"
    layout = workbook["Original table"]
    assert layout.max_row == 3 and layout.max_column == 3
    assert layout["A2"].value == "A1"
    assert layout["B2"].value == 12.5
    assert layout["C2"].value == "8.2 ± 0.4"
    assert layout.column_dimensions["B"].width > layout.column_dimensions["A"].width
    assert workbook["Audit"].max_row == 12

    with ZipFile(output) as archive:
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
        assert b"<autoFilter" not in sheet_xml


def test_explicit_numeric_region_preserves_measurements_inside_header_rows(tmp_path):
    from tablescan_local.domain import CellRuleRegion
    from tablescan_local.constraints import ValueConstraints
    job = make_job()
    job.template.cell_rules = [CellRuleRegion("values", "Measurements", 0, 0, 1, 2,
                                              ValueConstraints("numeric", 0, 100, 1))]
    for column, value in [(1, "57.3"), (2, "2.8")]:
        job.pages[0].cell(0, column).final_text = value
        job.template.column_rules[column].name = f"Column {column + 1}"
    workbook = load_workbook(export_job(job, tmp_path / "first-row.xlsx"), data_only=True)
    rows = list(workbook["Data"].values)
    assert [rows[1][6], rows[2][6]] == [57.3, 2.8]
    assert [rows[1][5], rows[2][5]] == ["2", "3"]
    assert workbook["Original table"]["B1"].value == 57.3
    assert job.template.header_rows == 1  # Export never mutates the saved template.
    workbook.close()


def test_numeric_headers_without_explicit_regions_remain_headers(tmp_path):
    workbook = load_workbook(export_job(make_job(), tmp_path / "headers.xlsx"), data_only=True)
    assert workbook["Data"].max_row == 4
    assert [row[6] for row in list(workbook["Data"].values)[1:]] == [12.5, "8.2 ± 0.4", "15%"]
    workbook.close()


def test_layout_sheet_keeps_excluded_measurements_empty(tmp_path):
    job = make_job()
    job.pages[0].excluded_rows = [2]
    job.pages[0].cell(2, 1).status = "excluded"
    job.pages[0].cell(2, 2).status = "excluded"

    workbook = load_workbook(export_job(job, tmp_path / "excluded.xlsx"), data_only=True)
    layout = workbook["Original table"]

    assert layout["A3"].value == "A2"
    assert layout["B3"].value is None and layout["C3"].value is None


def test_column_exclusion_overlaps_rows_and_restores_reviewed_values(tmp_path):
    job = make_job(); page = job.pages[0]
    cell = page.cell(1, 1); cell.final_text = '13.7'; cell.status = 'corrected'
    page.set_excluded('column', 1, True, job.template)
    page.set_excluded('row', 1, True, job.template)
    page.set_excluded('column', 1, False, job.template)
    assert cell.status == 'excluded' and cell.final_text == ''
    restored = JobResult.from_dict(job.to_dict())
    page = restored.pages[0]; cell = page.cell(1, 1)
    page.set_excluded('row', 1, False, restored.template)
    assert cell.final_text == '13.7' and cell.status == 'corrected'
    page.set_excluded('column', 1, True, restored.template)
    wb = load_workbook(export_job(restored, tmp_path/'columns.xlsx'))
    assert wb['Original table']['B1'].value == '1'
    assert wb['Original table']['B2'].value is None
    assert all(row[5] != 'Measure 1' for row in list(wb['Data'].values)[1:])


def test_compact_fields_include_all_pages_and_literal_metadata(tmp_path):
    job = make_job()
    second = PageResult.from_dict(job.pages[0].to_dict()); second.page_index = 1
    second.fields[0].final_text = '=1+1'
    second.fields[1].final_text = '0017'
    job.pages.append(second)
    wb = load_workbook(export_job(job, tmp_path/'compact-fields.xlsx', mode='compact'))
    assert wb.sheetnames == ['Original table', 'Original table 2', 'Fields']
    assert list(wb['Fields'].values)[1:] == [(1,'Test','PLANTAR'), (1,'Side','Left Hind'), (2,'Test','=1+1'), (2,'Side','0017')]
    assert wb['Fields']['C4'].data_type == 's'


def test_export_keeps_raw_and_alternatives_as_literal_text(tmp_path):
    job = make_job()
    item = job.pages[0].cell(1, 1)
    item.raw_text = "=1+1"
    item.alternatives = "125 | I2.5"
    output = export_job(job, tmp_path / "audit.xlsx")
    wb = load_workbook(output, data_only=False)
    rows = list(wb["Audit"].values)
    record = next(r for r in rows[1:] if r[3] == "R2C2")
    assert record[4] == "=1+1"
    assert record[-3] == "125 | I2.5"
    assert all(c.data_type != "f" for row in wb["Audit"] for c in row)


def test_export_blocks_unconfirmed_numeric_and_retains_first_data_row(tmp_path):
    import pytest
    job = make_job()
    job.template.header_rows = 0
    item = job.pages[0].cell(0, 1)
    item.raw_text = item.final_text = "34.4"
    item.status = "automatic"
    item.flags = ["low_confidence"]
    with pytest.raises(ValueError, match="need review"):
        export_job(job, tmp_path / "blocked.xlsx")
    item.status = "confirmed"
    wb = load_workbook(export_job(job, tmp_path / "confirmed.xlsx"))
    assert wb["Data"]["G2"].value == 34.4


def test_confirmed_value_overrides_hard_rule_on_export(tmp_path):
    import pytest
    job = make_job()
    job.template.column_rules[1].value_format = "numeric"
    job.template.column_rules[1].decimal_places = 1
    job.template.column_rules[1].maximum = 100
    item = job.pages[0].cell(1, 1)
    item.final_text = "125"
    item.applied_rule = "1 decimal place"
    wb = load_workbook(export_job(job, tmp_path / "manual-override.xlsx"))
    assert wb["Original table"]["B2"].value == 125
    wb.close()
    item.final_text = "12.5"
    output = export_job(job, tmp_path / "valid.xlsx")
    rows = list(load_workbook(output)["Audit"].values)
    record = next(row for row in rows[1:] if row[3] == "R2C2")
    assert record[-4] == "1 decimal place"


def test_export_records_writer_adaptation_evidence(tmp_path):
    job = make_job()
    item = job.pages[0].cell(1, 1)
    item.writer_suggestion = "12.5"
    item.writer_evidence = "page-local prototypes: 4 independent cells"
    output = export_job(job, tmp_path / "writer-audit.xlsx")
    rows = list(load_workbook(output)["Audit"].values)
    record = next(row for row in rows[1:] if row[3] == "R2C2")
    assert record[-2] == item.writer_suggestion
    assert record[-1] == item.writer_evidence


def test_compact_exports_original_layout_and_fields_for_every_page(tmp_path):
    job = make_job()
    second = PageResult.from_dict(job.pages[0].to_dict())
    second.page_index = 1
    job.pages.append(second)
    output = export_job(job, tmp_path / "compact.xlsx", mode="compact")
    wb = load_workbook(output, data_only=True)
    assert wb.sheetnames == ["Original table", "Original table 2", "Fields"]
    for sheet in (wb["Original table"], wb["Original table 2"]):
        assert sheet["B2"].value == 12.5
        assert sheet["C2"].value == "8.2 ± 0.4"
    wb.close()


def test_compact_still_enforces_review_and_preserves_text_and_exclusions(tmp_path):
    import pytest
    job = make_job()
    item = job.pages[0].cell(1, 1)
    item.status = "automatic"
    item.flags = ["low_confidence"]
    with pytest.raises(ValueError, match="need review"):
        export_job(job, tmp_path / "blocked-compact.xlsx", mode="compact")
    item.status = "confirmed"
    item.final_text = "=1+1"
    job.template.column_rules[1].value_format = "text"
    job.pages[0].excluded_rows = [2]
    for mode in ("compact", "extended"):
        wb = load_workbook(export_job(job, tmp_path / f"{mode}.xlsx", mode=mode), data_only=False)
        sheet = wb["Original table"]
        assert sheet["B2"].value == "=1+1" and sheet["B2"].data_type == "s"
        assert sheet["B3"].value is None and sheet["C3"].value is None
        wb.close()

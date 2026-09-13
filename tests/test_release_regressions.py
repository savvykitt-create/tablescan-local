"""User-visible regressions from the 0.7.13 release audit."""
import errno
import os
from copy import deepcopy

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import cv2
import numpy as np
import pytest
from openpyxl import Workbook, load_workbook
from PIL import Image
from PySide6.QtCore import Qt

from tablescan_local.constraints import ValueConstraints
from tablescan_local.domain import (
    CellResult, CellRuleRegion, FieldRegion, FieldResult, JobResult,
    NormalizedRect, PageResult, TableTemplate,
)
from tablescan_local.exporter import export_job
from tablescan_local.imaging import load_document
from tablescan_local.ui import create_application, QMessageBox, ReviewPage, TablePage, ValueRuleDialog, QDialog


@pytest.fixture
def warnings(monkeypatch):
    messages = []
    for method in ("warning", "critical", "information"):
        monkeypatch.setattr(QMessageBox, method, lambda *args: messages.append(args[1:]))
    return messages


@pytest.fixture
def window(qtbot, monkeypatch, tmp_path, warnings):
    monkeypatch.setenv("TABLESCAN_DATA_DIR", str(tmp_path / "data"))
    _, window = create_application()
    qtbot.addWidget(window)
    path = tmp_path / "original.png"
    cv2.imwrite(str(path), np.full((200, 300, 3), 255, np.uint8))
    window.open_source(str(path))
    return window


def reviewed_window(window):
    result = JobResult(window.source_path, deepcopy(window.template), [PageResult(
        0, window.source_path, [CellResult(0, 1, "1.25", "1.25", 1, status="confirmed")], [],
    )])
    window._recognition_done(result)
    return result


@pytest.mark.parametrize("failure", ["image", "detection", "draft", "cancel", "invalid_template"])
def test_failed_import_keeps_complete_active_context(window, tmp_path, monkeypatch, warnings, failure):
    result = reviewed_window(window)
    before = (window.job_id, window.source_path, window.stored_source_path, window.images,
              window.template, window.file_title.text())
    candidate = tmp_path / "next.png"
    cv2.imwrite(str(candidate), np.full((200, 300, 3), 255, np.uint8))
    if failure == "image":
        candidate.write_bytes(b"not a PNG")
    elif failure == "detection":
        monkeypatch.setattr(window, "_detect_or_manual", lambda *_: (_ for _ in ()).throw(RuntimeError("detection failed")))
    elif failure == "draft":
        monkeypatch.setattr(window.store, "save_draft", lambda *_: (_ for _ in ()).throw(OSError("write failed")))
    elif failure == "cancel":
        from tablescan_local.ui import TemplateChoiceDialog
        window.store.save_template(window.template)
        from types import SimpleNamespace
        monkeypatch.setattr("tablescan_local.ui.rank_templates", lambda *_: [SimpleNamespace(
            score=.9, template=window.template, reasons=[],
        )])
        monkeypatch.setattr(TemplateChoiceDialog, "exec", lambda *_: QDialog.DialogCode.Rejected)
    reused = deepcopy(window.template) if failure == "invalid_template" else None
    if reused:
        reused.row_guides = [0]
    window.open_source(str(candidate), reused)
    assert (window.job_id, window.source_path, window.stored_source_path) == before[:3]
    assert window.images is before[3] and window.template is before[4]
    assert window.file_title.text() == before[5] and window.result is result
    result.pages[0].cells[0].final_text = "2.50"
    window._result_changed(result)
    assert window.store.job_count() == 1
    assert len(list(window.store.jobs_dir.iterdir())) == 1
    assert window.store.load_job(window.job_id)[1].pages[0].cells[0].final_text == "2.50"


def test_bad_recent_draft_cannot_mix_an_old_result_with_new_images(window, tmp_path, warnings):
    result = reviewed_window(window)
    active_id, images = window.job_id, window.images
    other, _ = window.store.import_source(window.source_path)
    (window.store.jobs_dir / other / "working-template.json").write_text("{broken")
    window.open_recent(other)
    assert warnings and window.job_id == active_id and window.images is images
    assert window.result is result


def test_closing_unchanged_review_does_not_hide_saved_results(window, qtbot):
    original = reviewed_window(window)
    job_id = window.job_id
    assert window.close()
    _, reopened = create_application()
    qtbot.addWidget(reopened)
    reopened.open_recent(job_id)
    assert reopened.result.to_dict() == original.to_dict()
    assert reopened.stack.currentIndex() == 2


def table_job():
    template = TableTemplate("audit", "Audit", NormalizedRect(0, 0, 1, 1),
                             [0, .25, .5, .75, 1], [0, .33, .66, 1], header_rows=1)
    template.ensure_column_rules()
    template.cell_rules = [CellRuleRegion("m", "Measurements", 1, 3, 1, 2,
                                          ValueConstraints("numeric", -10, 100, 3))]
    rows = [["ID", "Value A", "Value B"], ["00123", "12.345", "0.006"],
            ["0000000000000000000124", "45.678", "-1.250"], ["00125", "90.123", "2.500"]]
    cells = [CellResult(r, c, text, text, 1, status="confirmed")
             for r, row in enumerate(rows) for c, text in enumerate(row)]
    return JobResult("synthetic.png", template, [PageResult(0, "synthetic.png", cells, [])])


@pytest.mark.parametrize("label_format,expected_numeric", [("complex_numeric", True), ("text", False)])
def test_label_ocr_keeps_configured_recognizer_while_review_preserves_identifier_text(tmp_path, monkeypatch, label_format, expected_numeric):
    from tablescan_local.pipeline import process_page
    from tablescan_local.ocr import OcrValue
    template = table_job().template
    template.column_rules[0].value_format = label_format
    calls = []
    class Engine:
        def recognize_cell(self, crop, numeric, constraints=None, retry_crops=None):
            calls.append(numeric)
            return OcrValue("00123", .99)
    monkeypatch.setattr("tablescan_local.pipeline.detect_crossed_rows", lambda *_: [])
    result = process_page(np.full((300, 400, 3), 255, np.uint8), "fixture.png", 0,
                          template, Engine(), crop_directory=tmp_path)
    assert calls[0] is False  # Header remains text even above a numeric label column.
    assert calls[3] is expected_numeric
    assert result.cell(1, 0).final_text == "00123"
    assert template.cell_constraints(1, 0)[0] is None


def review(qtbot, job):
    widget = ReviewPage()
    qtbot.addWidget(widget)
    widget.set_result([np.full((300, 400, 3), 255, np.uint8)], job)
    return widget


def test_header_single_bulk_and_export_share_text_semantics(qtbot, tmp_path, monkeypatch, warnings):
    job = table_job()
    header = job.pages[0].cell(0, 2)
    header.raw_text = header.final_text = "VALUEB"
    header.status, header.flags = "automatic", ["low_confidence"]
    widget = review(qtbot, job)
    widget._cell_clicked(0, 2)
    widget.correct_value.setText("Value B")
    widget.confirm_current()
    assert header.final_text == "Value B" and header.status == "corrected" and not warnings
    header.status = "automatic"
    monkeypatch.setattr(QMessageBox, "question", lambda *_: QMessageBox.StandardButton.Yes)
    widget.confirm_all_uncertain()
    assert header.status == "confirmed"
    wb = load_workbook(export_job(job, tmp_path / "headers.xlsx"))
    assert wb["Original table"]["C1"].value == "Value B"
    wb.close()


@pytest.mark.parametrize("mode", ["compact", "extended"])
def test_export_preserves_identifiers_and_declared_decimal_precision(tmp_path, mode):
    job = table_job()
    wb = load_workbook(export_job(job, tmp_path / f"{mode}.xlsx", mode=mode))
    sheet = wb["Original table"]
    assert sheet["A2"].value == "00123" and sheet["A2"].data_type == "s"
    assert sheet["A3"].value == "0000000000000000000124"
    assert sheet["B2"].value == 12.345 and sheet["B2"].number_format == "0.000"
    assert sheet["C2"].value == .006 and sheet["C2"].number_format == "0.000"
    assert sheet["C3"].value == -1.25 and sheet["C3"].number_format == "0.000"
    if mode == "extended":
        values = list(wb["Data"].values)
        assert values[1][2] == "00123"
        assert wb["Data"]["E2"].number_format == "0.000"
    wb.close()


@pytest.mark.parametrize("places,value", [(1, "12.3"), (2, "12.30"), (3, "0.006"), (8, "0.00000006")])
def test_export_number_format_follows_scale_with_or_without_a_rule(tmp_path, places, value):
    job = table_job()
    job.template.cell_rules.clear()
    for explicit in (False, True):
        job.template.column_rules[1].decimal_places = places if explicit else None
        job.template.column_rules[1].value_format = "numeric"
        for row in (1, 2, 3):
            job.pages[0].cell(row, 1).final_text = value
        wb = load_workbook(export_job(job, tmp_path / f"scale-{explicit}.xlsx", mode="compact"))
        assert wb.active["B2"].number_format == "0." + "0" * places
        assert wb.active["B2"].value == float(value)
        wb.close()


@pytest.mark.parametrize("kind,required,bad,good", [("text", True, "", "SAMPLE-1"),
                                                    ("integer", False, "abc", "123"),
                                                    ("numeric", False, "abc", "12.5")])
def test_fields_are_validated_on_single_bulk_and_export(qtbot, tmp_path, monkeypatch, warnings, kind, required, bad, good):
    job = table_job()
    region = FieldRegion("f", "Sample", NormalizedRect(0, 0, .2, .2), kind=kind, required=required)
    job.template.fields = [region]
    field = FieldResult("f", "Sample", good, good, 1, status="confirmed")
    job.pages[0].fields = [field]
    widget = review(qtbot, job)
    widget.current_kind, widget.current_index = "field", 0
    widget.correct_value.setText(bad)
    widget.confirm_current()
    assert field.final_text == good and warnings
    field.final_text, field.status, field.flags = bad, "automatic", ["low_confidence"]
    monkeypatch.setattr(QMessageBox, "question", lambda *_: QMessageBox.StandardButton.Yes)
    widget.confirm_all_uncertain()
    assert field.status == "automatic"
    field.status = "confirmed"  # Old/tampered saved data must not bypass export validation.
    with pytest.raises(ValueError, match="Sample"):
        export_job(job, tmp_path / "invalid.xlsx")
    assert not (tmp_path / "invalid.xlsx").exists()
    field.final_text = good
    assert export_job(job, tmp_path / "valid.xlsx").exists()


@pytest.mark.parametrize("failure", ["write", "validation"])
def test_failed_export_preserves_existing_file_and_removes_temporary_file(tmp_path, monkeypatch, failure):
    output = export_job(table_job(), tmp_path / "existing.xlsx")
    before = output.read_bytes()
    if failure == "write":
        def failed_save(_self, path):
            path.write_bytes(b"partial output")
            raise OSError(errno.ENOSPC, "No space left on device")
        monkeypatch.setattr(Workbook, "save", failed_save)
    else:
        monkeypatch.setattr("tablescan_local.exporter.load_workbook", lambda *_args, **_kwargs: (_ for _ in ()).throw(ValueError("invalid workbook")))
    with pytest.raises((OSError, ValueError)):
        export_job(table_job(), output)
    assert output.read_bytes() == before
    assert list(tmp_path.iterdir()) == [output]


def test_advanced_dialog_applies_changed_range(qtbot, monkeypatch):
    page = TablePage()
    qtbot.addWidget(page)
    job = table_job()
    job.template.cell_rules.clear()
    page.set_document(np.full((300, 400, 3), 255, np.uint8), job.template)
    page.canvas.select_region(0, 3, 1, 2)
    def accept(dialog):
        dialog.coordinates[0].setValue(2)
        dialog.selection = tuple(spin.value() - 1 for spin in dialog.coordinates)
        return QDialog.DialogCode.Accepted
    monkeypatch.setattr(ValueRuleDialog, "exec", accept)
    page._advanced_for_selection()
    assert job.template.cell_rules[-1].address() == "B2:C4"


def test_working_grid_autosaves_and_close_flushes_latest_edit(window, qtbot, monkeypatch):
    job_id = window.job_id
    window.table_page.template_name.setText("Autosaved before OCR")
    window.table_page.template_name.editingFinished.emit()
    qtbot.waitUntil(lambda: window.store.load_draft(job_id).name == "Autosaved before OCR")
    window.table_page.template_name.setText("Last edit before close")
    assert window.close()
    _, reopened = create_application()
    qtbot.addWidget(reopened)
    reopened.open_recent(job_id)
    assert reopened.template.name == "Last edit before close" and reopened.result is None
    assert reopened.stack.currentIndex() == 1


def test_failed_draft_save_prevents_closing_or_switching_documents(window, monkeypatch, warnings):
    window.table_page.template_name.setText("Must keep this edit")
    active_id = window.job_id
    monkeypatch.setattr(window.store, "save_draft", lambda *_: (_ for _ in ()).throw(OSError("disk full")))
    assert not window.close()
    window.open_source(window.source_path)
    assert window.job_id == active_id and window._pending_draft is not None
    assert warnings
    monkeypatch.undo()  # Restore writes before pytest-qt closes the window.


@pytest.mark.parametrize("orientation", range(1, 9))
def test_all_jpeg_exif_orientations(orientation, tmp_path):
    source = np.arange(8 * 12 * 3, dtype=np.uint8).reshape(8, 12, 3)
    path = tmp_path / "photo.jpg"
    exif = Image.Exif()
    exif[274] = orientation
    Image.fromarray(source).save(path, exif=exif)
    with Image.open(path) as image:
        pixels = np.asarray(image.convert("RGB"))[:, :, ::-1]
    expected = {1: lambda a: a, 2: lambda a: a[:, ::-1], 3: lambda a: a[::-1, ::-1],
                4: lambda a: a[::-1], 5: lambda a: a.transpose(1, 0, 2),
                6: lambda a: np.rot90(a, -1), 7: lambda a: a.transpose(1, 0, 2)[::-1, ::-1],
                8: lambda a: np.rot90(a, 1)}[orientation](pixels)
    assert np.array_equal(load_document(path)[0], expected)


def test_oldest_of_101_jobs_is_accessible_by_paging_and_search(window, tmp_path, qtbot):
    reviewed_window(window)
    original_id = window.job_id
    other = tmp_path / "other.png"
    cv2.imwrite(str(other), np.full((200, 300, 3), 255, np.uint8))
    for _ in range(100):
        window.store.import_source(other)
    window.store.connection.execute("UPDATE jobs SET updated_at='2000-01-01T00:00:00+00:00' WHERE id=?", (original_id,))
    window.store.connection.commit()
    window.refresh_recent()
    page = window.files_page
    assert page.recent.rowCount() == 20 and page.history_next.isEnabled()
    for _ in range(5):
        page.history_next.click()
    assert page.recent.rowCount() == 1 and not page.history_next.isEnabled()
    assert page.recent.item(0, 0).data(Qt.ItemDataRole.UserRole) == original_id
    page.history_search.setText("original")
    assert page.history_offset == 0 and page.recent.rowCount() == 1
    page.recent.doubleClicked.emit(page.recent.model().index(0, 0))
    assert window.job_id == original_id and window.result is not None

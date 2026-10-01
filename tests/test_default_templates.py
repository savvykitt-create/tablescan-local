import json
from pathlib import Path

from tablescan_local.storage import LocalStore


def test_first_launch_installs_four_templates_and_preserves_deletions(tmp_path):
    store = LocalStore(tmp_path)
    store.install_default_templates()
    templates = store.load_templates()
    assert {t.name for t in templates} == {'Von Frey', 'Plantar', 'Staircase', 'Staircase portrait (30 animals)'}
    assert all(t.auto_fit_rows and Path(t.reference_source_path).exists() for t in templates)
    assert all(Path(__file__).parents[1].joinpath('src/tablescan_local/default_templates', key + '.xlsx').exists()
               for key in ('von_frey', 'plantar', 'staircase', 'staircase_portrait'))
    removed = templates[0]
    store.delete_template(removed.id)
    store.install_default_templates()
    assert len(store.load_templates()) == 3
    assert store.job_count() == 0


def test_default_install_preserves_existing_family(tmp_path):
    store = LocalStore(tmp_path)
    store.install_default_templates()
    existing = store.load_templates()[0]
    existing.name = 'My edited form'
    store.save_template(existing)
    (tmp_path / '.default-templates-installed').unlink()
    store.install_default_templates()
    assert len(store.load_templates()) == 4
    assert any(t.name == 'My edited form' for t in store.load_templates())


def test_existing_install_gets_portrait_once_without_restoring_deleted_forms(tmp_path):
    store = LocalStore(tmp_path)
    (tmp_path / '.default-templates-installed').write_text('installed\n')
    store.install_default_templates()
    templates = store.load_templates()
    assert [t.name for t in templates] == ['Staircase portrait (30 animals)']
    store.delete_template(templates[0].id)
    store.install_default_templates()
    assert store.load_templates() == []


def test_portrait_printable_form_has_30_animals_and_matching_ocr_rules():
    import pypdfium2
    from openpyxl import load_workbook
    from tablescan_local.domain import TableTemplate
    from tablescan_local.imaging import detect_grid, load_document
    from tablescan_local.template_fit import fit_template
    root = Path(__file__).parents[1] / 'src/tablescan_local/default_templates'
    template = TableTemplate.from_dict(json.loads((root / 'staircase_portrait.json').read_text()))
    with pypdfium2.PdfDocument(root / 'staircase_portrait.pdf') as pdf:
        assert len(pdf) == 1
        w, h = pdf[0].get_size()
        assert w < h and abs(w / h - 210 / 297) < .001
    image = load_document(root / 'staircase_portrait.pdf')[0]
    detection = detect_grid(image)
    fitted = fit_template(template, detection, image.shape[1] / image.shape[0])
    from tablescan_local.template_matcher import rank_templates
    defaults = [TableTemplate.from_dict(json.loads(path.read_text())) for path in root.glob('*.json')]
    assert rank_templates(image.shape, detection, defaults)[0].template.id == template.id
    assert fitted.rows == 31 and fitted.columns == 18
    assert [f.name for f in fitted.fields if f.source != 'fixed'] == ['Date', 'Session', 'Operator']
    for col in range(1, 17):
        rule, _ = fitted.cell_constraints(30, col)
        assert rule.value_format == 'integer'
        assert rule.hard_errors('-1') and not rule.hard_errors('')
        assert fitted.fixed_value(0, col) == template.fixed_value(0, col)
    wb = load_workbook(root / 'staircase_portrait.xlsx')
    ws = wb.active
    assert ws.max_row == 36 and ws.max_column == 18
    assert ws.page_setup.orientation == 'portrait'
    assert ws.page_setup.fitToWidth == ws.page_setup.fitToHeight == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage
    assert 'A$1:$R$36' in str(ws.print_area)
    assert ws['A36'].value == '=ROW()-6'
    assert all(ws.cell(r, c).value is None for r in range(7, 37) for c in range(2, 19))
    wb.close()

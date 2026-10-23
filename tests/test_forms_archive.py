import importlib.util
from io import BytesIO
from pathlib import Path
import re
import shutil
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook


ROOT = Path(__file__).parents[1]
SOURCE = ROOT / 'src/tablescan_local/default_templates'
README = ROOT / 'examples/lab_forms_excel/README.md'
spec = importlib.util.spec_from_file_location('make_forms_archive', ROOT / 'packaging/make_forms_archive.py')
forms = importlib.util.module_from_spec(spec)
spec.loader.exec_module(forms)


@pytest.mark.parametrize('path', sorted(SOURCE.glob('*.xlsx')) +
                         sorted((ROOT / 'examples/lab_forms_excel/xlsx').glob('*.xlsx')),
                         ids=lambda path: f'{path.parent.name}/{path.name}')
def test_shipped_workbooks_have_excel_compatible_print_settings(path):
    forms.validate_workbook(path)


def test_archive_contains_readable_complete_forms(tmp_path):
    output = forms.make_archive(tmp_path / 'forms.zip', SOURCE, README)
    with ZipFile(output) as archive:
        assert archive.testzip() is None
        assert set(archive.namelist()) == {'README.md'} | {
            f'{folder}/{key}.{extension}' for key in forms.FORM_KEYS
            for folder, extension in [('Excel', 'xlsx'), ('PDF', 'pdf'), ('templates', 'json')]
        }
        for key in forms.FORM_KEYS:
            workbook = load_workbook(BytesIO(archive.read(f'Excel/{key}.xlsx')))
            sheet = workbook.active
            last_row = 37 if key == 'staircase_portrait' else 26
            header_row = 7 if key == 'staircase_portrait' else 6
            assert sheet.max_row == last_row
            assert sheet[f'A{header_row}'].value == 'ID'
            assert sheet[f'A{last_row}'].value == f'=ROW()-{header_row}'
            workbook.close()
            assert archive.read(f'PDF/{key}.pdf') == (SOURCE / f'{key}.pdf').read_bytes()


@pytest.mark.parametrize('existing', [False, True])
def test_rejects_excel_repair_trigger_before_writing_archive(tmp_path, existing):
    source = tmp_path / 'source'
    shutil.copytree(SOURCE, source)
    path = source / 'plantar.xlsx'
    with ZipFile(path) as workbook:
        parts = [(info, workbook.read(info)) for info in workbook.infolist()]
    with ZipFile(path, 'w') as workbook:
        for info, data in parts:
            if info.filename == 'xl/worksheets/sheet1.xml':
                options = re.search(rb'<printOptions\b[^>]*/>', data).group()
                data = data.replace(options, b'')
                setup = re.search(rb'<pageSetup\b[^>]*/>', data).group()
                data = data.replace(setup, setup + options)
            workbook.writestr(info, data)
    output = tmp_path / 'forms.zip'
    if existing:
        output.write_bytes(b'previous archive')
    with pytest.raises(ValueError, match='plantar.xlsx: invalid print settings order'):
        forms.make_archive(output, source, README)
    if existing:
        assert output.read_bytes() == b'previous archive'
    else:
        assert not output.exists()

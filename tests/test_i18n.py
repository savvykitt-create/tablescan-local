import ast
import copy
import json
import re
from pathlib import Path
from string import Formatter

import numpy as np
import pytest
from PySide6.QtCore import Qt
from PySide6.QtWidgets import QLabel as NativeLabel, QAbstractButton, QGroupBox, QComboBox as NativeCombo

from tablescan_local.constraints import ValueConstraints
from tablescan_local.domain import CellResult, FieldRegion, FieldResult, JobResult, NormalizedRect, PageResult, TableTemplate
from tablescan_local.i18n import SUPPORTED_LANGUAGES, _catalog, language, set_language, tr
from tablescan_local.rule_editor import ValueRuleDialog, optional_number
from tablescan_local.ui import create_application, ExportOptionsDialog, ReviewPage


@pytest.fixture(autouse=True)
def reset_language():
    set_language('en')
    yield
    set_language('en')


def test_every_message_has_three_complete_translations_with_same_placeholders():
    def placeholders(text):
        return sorted((field, spec, conversion) for _, field, spec, conversion in Formatter().parse(text) if field is not None)
    for source, versions in _catalog.items():
        assert set(versions) == {'en', 'cs', 'ru'}, source
        for code, translated in versions.items():
            assert translated.strip(), (source, code)
            assert placeholders(translated) == placeholders(source), (source, code)
            if code in ('en', 'cs'):
                assert not re.search('[А-Яа-яЁё]', translated), (source, code)


def test_no_unmarked_russian_ui_literals_or_missing_catalog_entries():
    root = Path(__file__).parents[1] / 'src/tablescan_local'
    for file in root.glob('*.py'):
        tree = ast.parse(file.read_text(encoding="utf-8"))
        parents = {child: node for node in ast.walk(tree) for child in ast.iter_child_nodes(node)}
        for node in ast.walk(tree):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id == 'tr':
                assert isinstance(node.args[0], ast.Constant), (file.name, node.lineno)
                assert node.args[0].value in _catalog, (file.name, node.lineno)
            if isinstance(node, ast.Constant) and isinstance(node.value, str) and re.search('[А-Яа-яЁё]', node.value):
                if node.value == 'Русский':
                    continue  # Native language name remains recognizable in every locale.
                parent = parents.get(node)
                assert isinstance(parent, ast.Call) and isinstance(parent.func, ast.Name) and parent.func.id == 'tr', (file.name, node.lineno)


def test_english_default_and_language_persistence(tmp_path, monkeypatch, qtbot):
    monkeypatch.setenv('TABLESCAN_DATA_DIR', str(tmp_path / 'app'))
    # System/application locale before startup must not change the default.
    set_language('ru')
    _, window = create_application()
    qtbot.addWidget(window)
    assert language() == 'en'
    assert window.language_select.currentData() == 'en'
    assert 'Settings' in window.section_buttons[2].text()
    window.language_select.setCurrentIndex(window.language_select.findData('cs'))
    assert language() == 'cs'
    assert 'Nastavení' in window.section_buttons[2].text()
    _, reopened = create_application()
    qtbot.addWidget(reopened)
    assert language() == 'cs' and reopened.language_select.currentData() == 'cs'
    assert reopened.review_page.confirm_button.text() == 'Potvrdit a pokračovat'


def test_language_switch_preserves_review_input_selection_image_and_result(qtbot, tmp_path, monkeypatch):
    monkeypatch.setenv('TABLESCAN_DATA_DIR', str(tmp_path / 'review'))
    _, window = create_application()
    qtbot.addWidget(window)
    template = TableTemplate('t', 'Мой шаблон', NormalizedRect(0, 0, 1, 1), [0, .5, 1], [0, .5, 1], row_label_columns=0)
    template.ensure_column_rules()
    template.fields = [FieldRegion('f', 'Моё поле', NormalizedRect(0, 0, .5, .5))]
    cell = CellResult(0, 0, '12.1', '12.1', .8, flags=['model_disagreement'], writer_suggestion='12.7', applied_rule='Старый аудит: Число')
    field = FieldResult('f', 'Моё поле', 'Русский исходный текст', 'Русский исходный текст', .99, status='confirmed')
    result = JobResult('source.png', template, [PageResult(0, 'source.png', [cell], [field])])
    review = window.review_page
    review.set_result([np.full((200, 200, 3), 255, np.uint8)], result)
    window.result = result
    window._navigate(2)
    window.show()
    review.details_toggle.setChecked(True)
    review.correct_value.setText('12,7')
    review.correct_value.setCursorPosition(2)
    snapshot = copy.deepcopy(result.to_dict())
    crop = review.crop_label._source_pixmap.cacheKey()
    for code, button, address in (('cs', 'Potvrdit a pokračovat', 'Buňka A1'), ('ru', 'Подтвердить и далее', 'Ячейка A1'), ('en', 'Confirm and continue', 'Cell A1')):
        window.language_select.setCurrentIndex(window.language_select.findData(code))
        assert review.correct_value.text() == '12,7'
        assert review.correct_value.cursorPosition() == 2
        assert review.current_index == 0 and review.canvas.active_cell == (0, 0)
        assert review.crop_label._source_pixmap.cacheKey() == crop
        assert result.to_dict() == snapshot
        assert review.confirm_button.text() == button
        assert review.address_label.text() == address
        assert address in review.crop_caption.text()
        assert '12.7' in review.writer_suggestion_button.text()
        assert review.details_toggle.isChecked()
        assert review.fields_table.item(0, 1).text() == 'Русский исходный текст'
        if code != 'ru':
            assert not re.search('[А-Яа-яЁё]', review.confidence_label.text())
    review.confirm_current()
    assert cell.final_text == '12.7'


def test_live_translation_in_open_rule_and_export_dialogs(qtbot):
    rule = ValueConstraints('numeric', 0, 100, 1, suggest_missing_decimal=True, require_decimal=True)
    dialog = ValueRuleDialog(rule)
    export = ExportOptionsDialog()
    qtbot.addWidget(dialog); qtbot.addWidget(export)
    dialog.minimum.setText('10.5')
    dialog.allowed.setText('12.5; 15.5')
    for code, apply, cancel in (('cs', 'Použít pravidlo', 'Zrušit'), ('ru', 'Применить правило', 'Отмена'), ('en', 'Apply rule', 'Cancel')):
        set_language(code)
        assert dialog.buttons.button(dialog.buttons.StandardButton.Save).text() == apply
        assert dialog.buttons.button(dialog.buttons.StandardButton.Cancel).text() == cancel
        assert dialog.minimum.text() == '10.5' and dialog.allowed.text() == '12.5; 15.5'
        assert dialog.kind.currentData() == 'numeric'
        assert dialog.places.value() == 1
        assert export.compact.text() == str(tr('Сокращённый — только исходная таблица'))


def test_edited_default_name_is_not_retranslated(qtbot, tmp_path, monkeypatch):
    monkeypatch.setenv('TABLESCAN_DATA_DIR', str(tmp_path / 'editor'))
    _, window = create_application()
    qtbot.addWidget(window)
    field = window.table_page.quick_rule_name
    field.setFocus()
    qtbot.keyClick(field, Qt.Key.Key_A, Qt.KeyboardModifier.ControlModifier)
    qtbot.keyClicks(field, 'My experiment')
    original = field.text()
    set_language('cs')
    assert field.text() == original


@pytest.mark.parametrize('code', ['en', 'cs', 'ru'])
def test_numeric_rules_and_export_values_are_locale_independent(code, tmp_path):
    from openpyxl import load_workbook
    from tablescan_local.exporter import export_job
    set_language(code)
    rule = ValueConstraints('numeric', 0, 100, 1)
    assert optional_number('12,5') == 12.5
    assert rule.hard_errors('12,5') == []
    assert rule.hard_errors('101.5') == ['above_maximum']
    template = TableTemplate('t', 'Text stays unchanged', NormalizedRect(0, 0, 1, 1), [0, 1], [0, 1], row_label_columns=0)
    template.ensure_column_rules()
    cell = CellResult(0, 0, '12,5', '12.5', .99, status='confirmed')
    job = JobResult('input.png', template, [PageResult(0, 'input.png', [cell], [])])
    path = export_job(job, tmp_path / f'{code}.xlsx', mode='compact')
    wb = load_workbook(path)
    assert wb.sheetnames == [str(tr('Исходная таблица'))]
    assert wb.active['A1'].value == 12.5
    wb.close()


def test_generated_names_are_frozen_as_document_data(qtbot):
    from PySide6.QtCore import QRectF
    from tablescan_local.ui import TablePage
    page = TablePage()
    qtbot.addWidget(page)
    template = TableTemplate('t', 'Untitled', NormalizedRect(0, 0, 1, 1), [0, .5, 1], [0, .5, 1])
    template.ensure_column_rules()
    page.set_document(np.full((200, 200, 3), 255, np.uint8), template)
    page._region_created('field', QRectF(10, 10, 50, 30))
    page.canvas.select_cells({(1, 1)})
    page._append_rule_regions(tr('Значения'), ValueConstraints('numeric'), '#336699')
    before = copy.deepcopy(template.to_dict())
    set_language('cs')
    assert template.to_dict() == before
    assert type(template.fields[0].name) is str
    assert type(template.cell_rules[0].name) is str


def test_standard_qt_buttons_follow_application_language(qtbot):
    from PySide6.QtWidgets import QDialogButtonBox
    for code, cancel, yes in (("en", "Cancel", "Yes"), ("cs", "Zrušit", "Ano"), ("ru", "Отмена", "Да")):
        set_language(code)
        box = QDialogButtonBox(QDialogButtonBox.StandardButton.Cancel | QDialogButtonBox.StandardButton.Yes)
        qtbot.addWidget(box)
        assert box.button(box.StandardButton.Cancel).text().replace("&", "") == cancel
        assert box.button(box.StandardButton.Yes).text().replace("&", "") == yes
